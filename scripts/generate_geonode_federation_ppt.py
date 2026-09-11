from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


OUT = Path("outputs")
OUT.mkdir(exist_ok=True)
PPTX_PATH = OUT / "GeoNexus-GeoNode-Federation-Collaboration.pptx"

CYAN = "06B6D4"
CYAN_DARK = "0891B2"
SLATE = "0F172A"
MUTED = "475569"
BG = "F8FAFC"
LINE = "CFFAFE"
WHITE = "FFFFFF"


def rgb(hex_value):
    return RGBColor.from_string(hex_value)


def add_title(slide, title, subtitle=None):
    box = slide.shapes.add_textbox(Inches(0.45), Inches(0.28), Inches(12.4), Inches(0.55))
    p = box.text_frame.paragraphs[0]
    p.text = title
    p.runs[0].font.size = Pt(25)
    p.runs[0].font.bold = True
    p.runs[0].font.color.rgb = rgb(SLATE)
    if subtitle:
        box2 = slide.shapes.add_textbox(Inches(0.48), Inches(0.78), Inches(12.1), Inches(0.35))
        p2 = box2.text_frame.paragraphs[0]
        p2.text = subtitle
        p2.runs[0].font.size = Pt(10)
        p2.runs[0].font.color.rgb = rgb(MUTED)


def add_capsule(slide, x, y, w, h, text, fill=WHITE, line=LINE, font=SLATE, size=12, bold=True):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line)
    shape.text_frame.clear()
    shape.text_frame.margin_left = Inches(0.16)
    shape.text_frame.margin_right = Inches(0.16)
    shape.text_frame.margin_top = Inches(0.06)
    shape.text_frame.margin_bottom = Inches(0.04)
    p = shape.text_frame.paragraphs[0]
    p.text = text
    p.alignment = PP_ALIGN.CENTER
    p.runs[0].font.size = Pt(size)
    p.runs[0].font.bold = bold
    p.runs[0].font.color.rgb = rgb(font)
    return shape


def add_text(slide, x, y, w, h, text, size=11, bold=False, color=MUTED):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    p.runs[0].font.size = Pt(size)
    p.runs[0].font.bold = bold
    p.runs[0].font.color.rgb = rgb(color)
    return box


def add_bullets(slide, x, y, w, h, bullets, size=10):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    for i, text in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = text
        p.level = 0
        p.runs[0].font.size = Pt(size)
        p.runs[0].font.color.rgb = rgb(MUTED)
        p.space_after = Pt(5)
    return box


def add_arrow(slide, x1, y1, x2, y2):
    line = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = rgb(CYAN_DARK)
    line.line.width = Pt(1.6)
    return line


def background(slide):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(BG)


def slide_cover(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background(slide)
    add_capsule(slide, 0.62, 0.46, 2.2, 0.42, "GeoNexus", fill=CYAN, line=CYAN, font=WHITE, size=13)
    add_text(slide, 0.7, 1.35, 11.8, 1.1, "GeoNode → GeoNode Federation → Collaborate", size=36, bold=True, color=SLATE)
    add_text(slide, 0.74, 2.52, 9.9, 0.7, "Cloud-native, hybrid-cloud and CAFE-inspired distributed computing architecture for a global GeoAI intelligence network.", size=16, color=MUTED)
    for i, label in enumerate(["GeoNode", "Federation", "Collaboration", "Global Intelligence"]):
        x = 0.85 + i * 3.0
        add_capsule(slide, x, 4.35, 2.25, 0.68, label, fill=WHITE, line=LINE, size=13)
        if i < 3:
            add_arrow(slide, x + 2.28, 4.69, x + 2.85, 4.69)
    add_text(slide, 0.8, 6.55, 11.8, 0.3, "Reference concept: CAFE collaborative analysis for distributed gridded environmental data, upgraded for GeoAI capabilities.", size=9, color=MUTED)


def slide_1(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background(slide)
    add_title(slide, "1. GeoNode: 主权可控的云原生地理智能能力节点", "GeoNode is not a data warehouse. It is Data + Compute + Model + Skill + Policy + Runtime.")
    center = add_capsule(slide, 4.7, 2.7, 3.9, 0.8, "GeoNode", fill=CYAN, line=CYAN, font=WHITE, size=20)
    items = [
        ("GeoData", 1.0, 1.35), ("GeoModel", 4.9, 1.05), ("GeoSkill", 8.8, 1.35),
        ("GeoAgent", 1.0, 4.95), ("GeoCompute", 4.9, 5.25), ("GeoPolicy", 8.8, 4.95),
    ]
    for label, x, y in items:
        add_capsule(slide, x, y, 2.6, 0.62, label)
        add_arrow(slide, x + 1.3, y + 0.62, 6.65, 3.1)
    add_bullets(slide, 0.75, 6.35, 11.8, 0.6, [
        "GeoNode keeps raw data local, exposes standardized capabilities, and returns derived outputs with provenance.",
        "It is deployable on public cloud, sovereign cloud, on-prem Kubernetes, edge K3s, or HPC clusters.",
    ], size=10)


def slide_2(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background(slide)
    add_title(slide, "2. GeoNode 内部技术栈", "A production GeoNode is a small cloud-native platform packaged for sovereign deployment.")
    layers = [
        ("Node Gateway", "Envoy / Kong / APISIX / GeoMCP endpoint"),
        ("Local Registry", "PostgreSQL + PostGIS + JSONB + pgvector + OpenSearch"),
        ("GeoData Layer", "STAC + COG + Zarr + GeoParquet + MinIO + TiTiler"),
        ("GeoCompute", "Kubernetes + Dask/Ray + Temporal Worker + GPU Operator"),
        ("GeoSkill Runtime", "FastAPI + Python + GDAL + Rasterio + GeoPandas + Xarray"),
        ("Policy & Trust", "OPA + Keycloak + SPIFFE/SPIRE + Vault + Sigstore"),
        ("Observability", "OpenTelemetry + Prometheus + Grafana + Loki"),
    ]
    y = 1.28
    for title, body in layers:
        add_capsule(slide, 0.75, y, 3.0, 0.48, title, fill=WHITE, line=LINE, size=11)
        add_text(slide, 4.05, y + 0.05, 8.4, 0.36, body, size=10, color=MUTED)
        y += 0.72


def slide_3(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background(slide)
    add_title(slide, "3. GeoData Layer: 面向 CAFE 的分布式格网环境数据", "Data locality and standardized metadata are the foundation of collaborative geospatial analysis.")
    grid = [
        ("COG", "Cloud Optimized GeoTIFF for raster imagery and flood maps"),
        ("Zarr / NetCDF", "Climate and gridded environmental time-series via Xarray/Dask"),
        ("GeoParquet", "Large vector and tabular geospatial analytics"),
        ("STAC", "Catalog metadata and search for satellite resources"),
        ("PostGIS", "Spatial indexing, coverage footprint, task regions"),
        ("MinIO/S3", "Object storage for artifacts and derived outputs"),
    ]
    for i, (title, body) in enumerate(grid):
        x = 0.7 + (i % 3) * 4.15
        y = 1.45 + (i // 3) * 2.0
        add_capsule(slide, x, y, 3.5, 0.58, title, fill=CYAN, line=CYAN, font=WHITE, size=14)
        add_text(slide, x + 0.15, y + 0.82, 3.2, 0.8, body, size=10)
    add_text(slide, 0.9, 6.15, 11.3, 0.5, "CAFE principle: keep distributed gridded environmental data near its owner; send computation and aggregate standardized derived results.", size=12, bold=True, color=CYAN_DARK)


def slide_4(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background(slide)
    add_title(slide, "4. GeoSkill Runtime: 从工具到领域能力", "GeoSkill packages a domain workflow with inputs, outputs, dependencies, runtime and policies.")
    steps = ["retrieve satellite", "segment water", "overlay boundaries", "zonal statistics", "generate report"]
    x = 0.8
    for i, step in enumerate(steps):
        add_capsule(slide, x, 2.5, 2.15, 0.62, step, fill=WHITE, line=LINE, size=10)
        if i < len(steps) - 1:
            add_arrow(slide, x + 2.18, 2.81, x + 2.58, 2.81)
        x += 2.45
    add_bullets(slide, 0.9, 4.1, 5.8, 1.6, [
        "Runtime: FastAPI + Python + Temporal Worker",
        "Libraries: GDAL, Rasterio, GeoPandas, Xarray, Dask/Ray, PyTorch/ONNX",
        "Deployment: Docker image scheduled inside GeoNode Kubernetes",
    ])
    add_bullets(slide, 7.0, 4.1, 5.5, 1.6, [
        "Inputs and outputs are declared in a GeoSkill manifest",
        "Policy can require derived-only outputs",
        "Execution trace is written back to Registry and GeoKG",
    ])


def slide_5(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background(slide)
    add_title(slide, "5. GeoNode Federation: 可信能力的协同网络", "Federation connects distributed GeoNodes without centralizing raw data.")
    center = add_capsule(slide, 4.75, 3.0, 3.8, 0.75, "GeoNode Federation", fill=CYAN, line=CYAN, font=WHITE, size=16)
    nodes = [("Mekong Node", 0.9, 1.2), ("Nairobi Node", 9.7, 1.2), ("Punjab Node", 0.9, 5.4), ("Singapore Compute", 9.4, 5.4)]
    for label, x, y in nodes:
        add_capsule(slide, x, y, 2.8, 0.64, label)
        add_arrow(slide, x + 1.4, y + 0.64, 6.65, 3.38)
    add_bullets(slide, 4.55, 4.3, 4.35, 1.15, [
        "Node discovery",
        "Capability aggregation",
        "Trust negotiation",
        "Policy matching",
        "Cross-node routing",
    ], size=10)


def slide_6(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background(slide)
    add_title(slide, "6. Federation Control Plane 技术栈", "Spring Boot is preferred for the governance-heavy enterprise control plane.")
    services = [
        "Node Registry Service", "Capability Federation Service", "Trust Service", "Policy Service",
        "Routing Service", "Task Coordination Service", "Audit Service", "Marketplace Service",
    ]
    for i, service in enumerate(services):
        x = 0.75 + (i % 4) * 3.1
        y = 1.35 + (i // 4) * 1.35
        add_capsule(slide, x, y, 2.55, 0.62, service, fill=WHITE, line=LINE, size=10)
    add_capsule(slide, 0.95, 4.75, 2.9, 0.62, "PostGIS + pgvector")
    add_capsule(slide, 4.05, 4.75, 2.4, 0.62, "Neo4j GeoKG")
    add_capsule(slide, 6.65, 4.75, 2.4, 0.62, "OpenSearch")
    add_capsule(slide, 9.25, 4.75, 2.4, 0.62, "Temporal")
    add_text(slide, 0.95, 6.05, 11.4, 0.48, "The control plane stores metadata, not raw sovereign datasets. It coordinates, routes and audits federated capabilities.", size=12, bold=True, color=CYAN_DARK)


def slide_7(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background(slide)
    add_title(slide, "7. Federation Router: 全球任务如何选节点", "Routing combines semantic match, spatial coverage, policy fit, trust, latency and cost.")
    formula = [
        ("Capability Match", "30%"), ("Spatial Coverage", "20%"), ("Policy Fit", "20%"),
        ("Trust Score", "15%"), ("Latency", "10%"), ("Cost", "5%"),
    ]
    for i, (name, score) in enumerate(formula):
        x = 0.75 + i * 2.05
        add_capsule(slide, x, 1.6, 1.68, 0.75, score, fill=CYAN, line=CYAN, font=WHITE, size=18)
        add_text(slide, x, 2.48, 1.7, 0.5, name, size=9, bold=True, color=SLATE)
    add_bullets(slide, 0.95, 3.55, 5.7, 1.8, [
        "PostGIS filters nodes by spatial coverage and jurisdiction.",
        "pgvector/OpenSearch finds semantically relevant capabilities.",
        "Neo4j traverses dependencies and compatibility relationships.",
    ])
    add_bullets(slide, 7.0, 3.55, 5.5, 1.8, [
        "OPA evaluates sovereignty and derived-only policies.",
        "Trust service ranks node reliability and validation history.",
        "Temporal starts the selected cross-node workflow.",
    ])


def slide_8(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background(slide)
    add_title(slide, "8. CAFE → GeoNexus: 协同分析如何升级为能力协作", "CAFE's distributed gridded data analysis becomes federated GeoAI capability orchestration.")
    left = ["Distributed gridded environmental data", "Local computation", "Standardized derived results", "Collaborative analysis"]
    right = ["Distributed GeoCapabilities", "GeoSkill near-data execution", "Provenance-rich outputs", "Collaborative intelligence network"]
    add_capsule(slide, 1.0, 1.25, 4.4, 0.58, "CAFE principle", fill=CYAN, line=CYAN, font=WHITE, size=14)
    add_capsule(slide, 7.8, 1.25, 4.4, 0.58, "GeoNexus upgrade", fill=CYAN, line=CYAN, font=WHITE, size=14)
    for i in range(4):
        y = 2.15 + i * 0.95
        add_capsule(slide, 0.9, y, 4.65, 0.55, left[i], fill=WHITE, line=LINE, size=10)
        add_arrow(slide, 5.75, y + 0.28, 7.45, y + 0.28)
        add_capsule(slide, 7.65, y, 4.8, 0.55, right[i], fill=WHITE, line=LINE, size=10)
    add_text(slide, 1.0, 6.25, 11.2, 0.45, "Global analysis = local compute + global aggregation + reusable skill intelligence.", size=14, bold=True, color=CYAN_DARK)


def slide_9(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background(slide)
    add_title(slide, "9. 混合云部署：Public Cloud + Sovereign Cloud + On-prem + Edge + HPC", "GeoNexus must meet UN, national government and developing-country infrastructure constraints.")
    clouds = [
        ("Public Cloud", "global registry, marketplace, dashboards"),
        ("Sovereign Cloud", "regulated national workflows"),
        ("On-prem", "mapping agencies, universities, labs"),
        ("Edge", "field deployment and low-network regions"),
        ("HPC", "climate grids, CMIP/ERA5, long simulations"),
    ]
    for i, (title, body) in enumerate(clouds):
        x = 0.75 + (i % 3) * 4.05
        y = 1.45 + (i // 3) * 2.05
        add_capsule(slide, x, y, 3.35, 0.62, title, fill=WHITE, line=LINE, size=13)
        add_text(slide, x + 0.12, y + 0.8, 3.1, 0.65, body, size=10)
    add_capsule(slide, 4.9, 5.7, 3.5, 0.7, "One Federation, Many Clouds", fill=CYAN, line=CYAN, font=WHITE, size=15)


def slide_10(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background(slide)
    add_title(slide, "10. 全球智能网络如何形成", "GeoNode owns local capability. Federation connects trusted capability. GeoNexus coordinates global intelligence.")
    stages = [
        ("1. Node", "institutions deploy GeoNodes"),
        ("2. Federation", "nodes publish capability manifests"),
        ("3. Collaborate", "tasks are decomposed and routed"),
        ("4. Learning Network", "results update registry, trust and GeoKG"),
    ]
    x = 0.85
    for i, (title, body) in enumerate(stages):
        add_capsule(slide, x, 2.0, 2.55, 0.7, title, fill=CYAN if i == 3 else WHITE, line=CYAN if i == 3 else LINE, font=WHITE if i == 3 else SLATE, size=13)
        add_text(slide, x + 0.05, 2.9, 2.45, 0.75, body, size=10)
        if i < 3:
            add_arrow(slide, x + 2.6, 2.35, x + 3.05, 2.35)
        x += 3.05
    add_bullets(slide, 1.0, 4.6, 11.4, 1.2, [
        "Each task produces provenance, model performance, node reliability metrics and reusable workflow templates.",
        "These signals are written back to Capability Registry, GeoKG, Trust Service and Skill Marketplace.",
        "The network becomes smarter as more nodes contribute capabilities and validation results.",
    ], size=11)


def slide_summary(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background(slide)
    add_title(slide, "Final Architecture Motto", "The GeoNexus federation loop")
    add_capsule(slide, 1.0, 2.2, 3.0, 0.75, "GeoNode", fill=WHITE, line=LINE, size=18)
    add_arrow(slide, 4.2, 2.58, 4.9, 2.58)
    add_capsule(slide, 5.05, 2.2, 3.35, 0.75, "GeoNode Federation", fill=WHITE, line=LINE, size=18)
    add_arrow(slide, 8.6, 2.58, 9.25, 2.58)
    add_capsule(slide, 9.4, 2.2, 2.75, 0.75, "Collaborate", fill=CYAN, line=CYAN, font=WHITE, size=18)
    add_text(slide, 1.1, 4.0, 11.1, 0.8, "GeoNode 负责本地能力。GeoNode Federation 连接可信能力。GeoNexus 协调全球智能。", size=22, bold=True, color=SLATE)
    add_text(slide, 1.15, 5.2, 10.8, 0.55, "The long-term moat is not only data or models, but a sovereign, reusable, federated GeoAI capability standard.", size=14, color=MUTED)


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    for fn in [slide_cover, slide_1, slide_2, slide_3, slide_4, slide_5, slide_6, slide_7, slide_8, slide_9, slide_10, slide_summary]:
        fn(prs)
    prs.save(PPTX_PATH)
    print(PPTX_PATH)


if __name__ == "__main__":
    build()
