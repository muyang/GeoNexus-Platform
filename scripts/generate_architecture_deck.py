from pathlib import Path

from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, A4
from reportlab.lib.units import inch
from reportlab.pdfgen import canvas

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


OUT = Path("outputs")
OUT.mkdir(exist_ok=True)

PDF_PATH = OUT / "GeoNexus-Architecture.pdf"
PPTX_PATH = OUT / "GeoNexus-Architecture.pptx"

CYAN = "06B6D4"
CYAN_DARK = "0891B2"
SLATE = "0F172A"
MUTED = "475569"
BG = "F8FAFC"
LINE = "CFFAFE"


layers = [
    ("Frontend Console", "Next.js + TypeScript + Tailwind + React Flow + MapLibre + deck.gl", "Zero-code model builder, DAG canvas, map visualization, task monitoring"),
    ("API / Control Plane", "Spring Boot + Spring Security + Spring Cloud Gateway", "Tenants, users, RBAC/ABAC, registry, tasks, audit, marketplace"),
    ("Workflow Plane", "Temporal + Redis", "Long-running jobs, retries, checkpoints, resume, event history"),
    ("GeoAI Execution Plane", "FastAPI + Python GeoAI Runtime", "GeoSkill runtime, MCP tool server, Agent runtime, raster/model processing"),
    ("Data & Knowledge Plane", "PostgreSQL/PostGIS + pgvector + OpenSearch + Neo4j", "Capability registry, spatial metadata, semantic search, GeoKG relationships"),
    ("Federated GeoNode Plane", "Kubernetes + Docker + MinIO/S3 + OPA + Keycloak", "Sovereign node execution, local data connectors, policy enforcement"),
    ("Observability & Governance", "OpenTelemetry + Prometheus + Grafana + Loki + Vault", "Tracing, logs, metrics, secrets, auditability, compliance"),
]

flows = [
    "1. User submits GeoTask in natural language",
    "2. Spring Boot validates tenant, quota, policy and creates task",
    "3. Temporal starts durable workflow",
    "4. FastAPI GeoAgent queries Registry + GeoKG through GeoMCP",
    "5. GeoSkill runs near data on eligible GeoNode",
    "6. Derived artifacts are stored in MinIO/S3 and indexed in PostGIS/Registry",
    "7. Frontend receives progress/events and shows traceable outputs",
]

choices = [
    ("Why Spring Boot", "Best for enterprise control plane: tenancy, security, audit, workflow APIs, registry governance and marketplace lifecycle."),
    ("Why FastAPI", "Best for GeoAI execution: Python GIS, raster processing, PyTorch/ONNX inference, LangGraph agents and MCP tools."),
    ("Why Temporal", "Durable workflow engine for geospatial long-running tasks: retries, checkpoints, cancellation and human-in-the-loop."),
    ("Why PostGIS", "The operational source of truth for spatial metadata, task regions, coverage footprints and capability search filters."),
    ("Why GeoKG", "Makes dependencies explicit: Skill uses model, requires data, runs on node, governed by policy, orchestrated by agent."),
]


def hex_color(value):
    return colors.HexColor(f"#{value}")


def draw_capsule(c, x, y, w, h, title, subtitle, body, fill="#FFFFFF"):
    c.setFillColor(colors.HexColor(fill))
    c.setStrokeColor(hex_color(LINE))
    c.roundRect(x, y, w, h, h / 2, fill=1, stroke=1)
    c.setFillColor(hex_color(SLATE))
    c.setFont("Helvetica-Bold", 11)
    c.drawString(x + 18, y + h - 22, title)
    c.setFillColor(hex_color(CYAN_DARK))
    c.setFont("Helvetica-Bold", 8)
    c.drawString(x + 18, y + h - 36, subtitle)
    c.setFillColor(hex_color(MUTED))
    c.setFont("Helvetica", 8)
    text = c.beginText(x + 18, y + h - 51)
    for line in wrap(body, 68):
        text.textLine(line)
    c.drawText(text)


def wrap(text, width):
    words = text.split()
    lines = []
    line = []
    size = 0
    for word in words:
        if size + len(word) + 1 > width:
            lines.append(" ".join(line))
            line = [word]
            size = len(word)
        else:
            line.append(word)
            size += len(word) + 1
    if line:
        lines.append(" ".join(line))
    return lines


def make_pdf():
    c = canvas.Canvas(str(PDF_PATH), pagesize=landscape(A4))
    width, height = landscape(A4)
    c.setFillColor(hex_color(BG))
    c.rect(0, 0, width, height, fill=1, stroke=0)

    c.setFillColor(hex_color(SLATE))
    c.setFont("Helvetica-Bold", 26)
    c.drawString(0.55 * inch, height - 0.65 * inch, "GeoNexus Serious Product Architecture")
    c.setFillColor(hex_color(MUTED))
    c.setFont("Helvetica", 10)
    c.drawString(0.55 * inch, height - 0.9 * inch, "Spring Boot control plane + FastAPI GeoAI execution plane + Temporal workflow + PostGIS/GeoKG federation")

    x = 0.55 * inch
    y = height - 1.55 * inch
    box_w = width - 1.1 * inch
    box_h = 0.58 * inch
    gap = 0.13 * inch
    for title, tech, body in layers:
        draw_capsule(c, x, y, box_w, box_h, title, tech, body)
        y -= box_h + gap

    c.showPage()
    c.setFillColor(hex_color(BG))
    c.rect(0, 0, width, height, fill=1, stroke=0)
    c.setFillColor(hex_color(SLATE))
    c.setFont("Helvetica-Bold", 24)
    c.drawString(0.55 * inch, height - 0.65 * inch, "Execution Flow and Technology Rationale")

    left_x = 0.65 * inch
    top_y = height - 1.2 * inch
    c.setFont("Helvetica-Bold", 14)
    c.setFillColor(hex_color(CYAN_DARK))
    c.drawString(left_x, top_y, "GeoTask Execution Flow")
    c.setFont("Helvetica", 10)
    c.setFillColor(hex_color(SLATE))
    text = c.beginText(left_x, top_y - 24)
    for item in flows:
        text.textLine(item)
        text.textLine("")
    c.drawText(text)

    right_x = width / 2 + 0.2 * inch
    c.setFont("Helvetica-Bold", 14)
    c.setFillColor(hex_color(CYAN_DARK))
    c.drawString(right_x, top_y, "Technology Selection Notes")
    y = top_y - 0.45 * inch
    for title, body in choices:
        draw_capsule(c, right_x, y - 0.42 * inch, width / 2 - 0.85 * inch, 0.52 * inch, title, "Selection rationale", body)
        y -= 0.65 * inch

    c.save()


def add_textbox(slide, x, y, w, h, text, font_size=14, bold=False, color=SLATE, align=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    p = frame.paragraphs[0]
    p.text = text
    if align:
        p.alignment = align
    run = p.runs[0]
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    return shape


def add_capsule(slide, x, y, w, h, title, subtitle, body):
    shape = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
    shape.line.color.rgb = RGBColor.from_string(LINE)
    frame = shape.text_frame
    frame.margin_left = Inches(0.22)
    frame.margin_right = Inches(0.22)
    frame.margin_top = Inches(0.1)
    frame.clear()
    p = frame.paragraphs[0]
    p.text = title
    p.runs[0].font.bold = True
    p.runs[0].font.size = Pt(12)
    p.runs[0].font.color.rgb = RGBColor.from_string(SLATE)
    p2 = frame.add_paragraph()
    p2.text = subtitle
    p2.runs[0].font.bold = True
    p2.runs[0].font.size = Pt(8)
    p2.runs[0].font.color.rgb = RGBColor.from_string(CYAN_DARK)
    p3 = frame.add_paragraph()
    p3.text = body
    p3.runs[0].font.size = Pt(8)
    p3.runs[0].font.color.rgb = RGBColor.from_string(MUTED)


def make_pptx():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(BG)
    add_textbox(slide, 0.55, 0.28, 12, 0.45, "GeoNexus Serious Product Architecture", 28, True)
    add_textbox(slide, 0.58, 0.75, 12, 0.28, "Spring Boot control plane + FastAPI execution plane + Temporal workflow + PostGIS/GeoKG federation", 11, False, MUTED)
    y = 1.25
    for title, tech, body in layers:
        add_capsule(slide, 0.55, y, 12.25, 0.68, title, tech, body)
        y += 0.82

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(BG)
    add_textbox(slide, 0.55, 0.32, 12, 0.45, "Execution Flow and Technology Rationale", 26, True)
    add_textbox(slide, 0.65, 1.1, 5.8, 0.35, "GeoTask Execution Flow", 15, True, CYAN_DARK)
    add_textbox(slide, 0.68, 1.55, 5.8, 4.8, "\n\n".join(flows), 11, False, SLATE)
    add_textbox(slide, 6.85, 1.1, 5.8, 0.35, "Technology Selection Notes", 15, True, CYAN_DARK)
    y = 1.55
    for title, body in choices:
        add_capsule(slide, 6.85, y, 5.85, 0.72, title, "Selection rationale", body)
        y += 0.86

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(BG)
    add_textbox(slide, 0.55, 0.32, 12, 0.45, "Recommended Service Boundary", 26, True)
    services = [
        ("Spring Boot", "Platform API, identity, tenants, registry governance, task service, audit, marketplace"),
        ("FastAPI", "GeoSkill runtime, MCP tool server, Agent runtime, Python GIS/model services"),
        ("Temporal", "Workflow orchestration, long-running jobs, retries, checkpoints, resume"),
        ("PostGIS + Neo4j", "Spatial metadata, capability registry, GeoKG relationships, search"),
        ("GeoNode Runtime", "Kubernetes worker packages deployed near sovereign data"),
        ("Observability", "OpenTelemetry traces, Prometheus metrics, Grafana dashboards, Loki logs"),
    ]
    x_positions = [0.65, 4.75, 8.85]
    y_positions = [1.35, 3.55]
    idx = 0
    for y in y_positions:
        for x in x_positions:
            title, body = services[idx]
            add_capsule(slide, x, y, 3.65, 1.35, title, "Service boundary", body)
            idx += 1

    prs.save(PPTX_PATH)


if __name__ == "__main__":
    make_pdf()
    make_pptx()
    print(PDF_PATH)
    print(PPTX_PATH)
