"""生成 2027 课题 1 技术路线图 PPTX（矢量形状 + 可编辑文本）。

运行方式：
    /tmp/docxenv/bin/python gen_pptx.py
输出：
    ../2027年度课题1技术路线图-矢量版.pptx
"""

import os
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

BASE = Path(__file__).resolve().parent
OUT = BASE.parent / "2027年度课题1技术路线图-矢量版.pptx"

# ── 配色（与 matplotlib 版一致） ──
C = {
    "title_bg":   RGBColor(0x2E, 0x5E, 0x8C),  # 深蓝（标题栏）
    "base":       RGBColor(0xB3, 0xD9, 0xF7),  # 浅蓝（底座/阶段）
    "core":       RGBColor(0xFF, 0xE6, 0x99),  # 浅黄（核心）
    "app":        RGBColor(0xC6, 0xE6, 0xC6),  # 浅绿（应用）
    "arch":       RGBColor(0xF4, 0xB9, 0xB9),  # 浅红（成果/输出）
    "phase":      RGBColor(0xE0, 0xE0, 0xE0),  # 浅灰（阶段框）
    "method":     RGBColor(0xE8, 0xE3, 0xF5),  # 浅紫（方法）
    "white":      RGBColor(0xFF, 0xFF, 0xFF),
    "txt":        RGBColor(0x1F, 0x29, 0x37),
    "muted":      RGBColor(0x6B, 0x72, 0x80),
    "edge_base":  RGBColor(0x3A, 0x6E, 0x9C),
    "edge_green": RGBColor(0x3E, 0x6B, 0x3E),
    "edge_red":   RGBColor(0x9A, 0x3B, 0x3B),
    "edge_gray":  RGBColor(0x7A, 0x7A, 0x7A),
}

FONT = "Hiragino Sans GB"


def _set_cjk(run, font_name=FONT):
    """设置东亚字体（python-pptx 默认不处理 CJK）。"""
    run.font.name = font_name
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = etree.SubElement(rPr, qn('a:ea'))
    ea.set('typeface', font_name)


def add_title_bar(slide, text, left, top, width, height, font_size=Pt(12)):
    """深蓝标题栏。"""
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    s.fill.solid()
    s.fill.fore_color.rgb = C["title_bg"]
    s.line.fill.background()
    tf = s.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.15)
    tf.margin_right = Inches(0.15)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.size = font_size
    run.font.color.rgb = C["white"]
    run.font.bold = True
    _set_cjk(run)
    return s


def add_box(slide, text, left, top, width, height, bg_color=C["white"],
            edge_color=C["edge_base"], font_size=Pt(9), bold_title="",
            round_corner=True, text_color=C["txt"]):
    """圆角矩形节点框。"""
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if round_corner else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shape_type, left, top, width, height)
    s.fill.solid()
    s.fill.fore_color.rgb = bg_color
    s.line.color.rgb = edge_color
    s.line.width = Pt(1.2)
    tf = s.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.04)

    lines = text.split("\n")
    for i, line in enumerate(lines):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = line
        if i == 0 and bold_title and len(lines) > 1:
            run.font.size = font_size + Pt(1)
            run.font.bold = True
        else:
            run.font.size = font_size
        run.font.color.rgb = text_color
        _set_cjk(run)
    return s


def add_arrow(slide, left, top, width, height, color=C["edge_base"],
              direction="right"):
    """连接箭头（RIGHT_ARROW 或 DOWN_ARROW）。"""
    shape_type = MSO_SHAPE.RIGHT_ARROW if direction == "right" else MSO_SHAPE.DOWN_ARROW
    s = slide.shapes.add_shape(shape_type, left, top, width, height)
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    return s


def add_textbox(slide, text, left, top, width, height, font_size=Pt(9),
                color=C["txt"], alignment=PP_ALIGN.CENTER, bold=False):
    """自由文本。"""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = alignment
    run = p.add_run()
    run.text = text
    run.font.size = font_size
    run.font.color.rgb = color
    run.font.bold = bold
    _set_cjk(run)
    return txBox


def _in(val):
    """简写：英寸常量"""
    return Inches(val)


# ═══════════════════════════════════════════════════════════════
# 幻灯片 1：项目总体技术路线图
# ═══════════════════════════════════════════════════════════════
def slide_overall(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    W, H = 13.33, 7.5

    # 标题栏
    add_title_bar(slide, "总体目标：开放、安全、普惠的全球地理信息智能框架（GeoNexus）——支撑 UN-GGKIC 全球创新枢纽",
                  _in(0.4), _in(0.15), _in(12.5), _in(0.5), Pt(11))

    # 课题1 核心底座
    add_box(slide, "课题1 核心底座：全球地理信息智能框架研发与集成\nGeoMCP 协议栈 · GAKG 合约注册中心 · 零信任安全合规 · CAFE 计算下推 · 低代码智能开发",
            _in(0.4), _in(0.85), _in(12.5), _in(0.8), C["core"], C["edge_base"], Pt(9.5))

    # 三大应用
    apps = [
        ("课题2：地球观测卫星\n智能服务系统", C["app"], C["edge_green"]),
        ("课题3：全球可持续发展\n评价服务系统", C["app"], C["edge_green"]),
        ("课题4：地理信息知识与\n远程学习系统", C["app"], C["edge_green"]),
    ]
    for i, (txt, bg, ec) in enumerate(apps):
        x = 0.4 + i * 4.25
        add_box(slide, txt, _in(x), _in(1.95), _in(3.85), _in(0.85), bg, ec, Pt(10))
        add_arrow(slide, _in(2.6 + i * 4.25), _in(1.72), _in(0.35), _in(0.16), C["edge_green"])

    # 能力贯通
    add_box(slide, "能力贯通：资源接入 → 语义对齐 → 智能调度 → 计算下推 → 决策交付（全链路自动化）",
            _in(0.4), _in(3.1), _in(12.5), _in(0.55), C["method"], C["edge_base"], Pt(9.5))

    # 三阶段
    phases = [
        ("2026 基础构建期", "GeoMCP V1.0 · GAKG v0.5\nCAFE 原型 · 前端 V1.0\n卫星索引 ≥100 颗", C["phase"]),
        ("2027 深化完善期\n(Beta)", "协议 V2.0 · 框架系统 V2.0\n伦理规范 · 合规引擎\nModel Builder ≥50 节点", C["base"]),
        ("2028 集成收尾期\n(业务化)", "β→业务化 · ≥400 卫星\nGeoKG ≥100 万实体\nUN-GGKIC 承载 · 15 国示范", C["phase"]),
    ]
    for i, (title_txt, detail, bg) in enumerate(phases):
        x = 0.4 + i * 4.25
        add_box(slide, f"{title_txt}\n{detail}", _in(x), _in(3.9), _in(3.85), _in(1.25),
                bg, C["edge_gray"], Pt(8.5))
        if i < 2:
            add_arrow(slide, _in(x + 3.85), _in(4.45), _in(0.4), _in(0.15), C["edge_gray"])

    # 考核指标
    add_box(slide,
            "2027 考核：合规引擎并发判定<50ms · 空间推理+10% · Model Builder ≥50 节点 · GeoNow 2027\n"
            "论文 3 篇 · 发明专利 3 项 · 软件著作权 2 项 · 项目中期成果报告",
            _in(0.4), _in(5.55), _in(12.5), _in(0.65), C["arch"], C["edge_red"], Pt(9))


# ═══════════════════════════════════════════════════════════════
# 幻灯片 2：课题1 总体技术路线图
# ═══════════════════════════════════════════════════════════════
def slide_subject1(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    add_title_bar(slide, "课题1：全球地理信息智能框架研发与集成——\"标准引领→安全筑基→引擎驱动→生态聚合\"",
                  _in(0.4), _in(0.15), _in(12.5), _in(0.5), Pt(11))

    # 四阶段递进
    stages = [
        ("标准引领\n[交互标准化]\nGeoMCP V2 协议栈\nGAKG 合约元模型", C["core"]),
        ("安全筑基\n[安全·合规协作]\n零信任网关\n跨域合规·水印", C["arch"]),
        ("引擎驱动\n[CAFE 下推·弹性]\nCentral+Worker 架构\nServerless 毫秒级伸缩", C["core"]),
        ("生态聚合\n[低代码·共享]\n可视化编排器\n预训练智能体矩阵", C["app"]),
    ]
    for i, (txt, bg) in enumerate(stages):
        x = 0.4 + i * 3.25
        add_box(slide, txt, _in(x), _in(0.85), _in(2.85), _in(1.45), bg, C["edge_base"], Pt(8.5))
        if i < 3:
            add_arrow(slide, _in(x + 2.85), _in(1.5), _in(0.4), _in(0.15), C["edge_base"])

    # 2027 纵向深化
    add_box(slide, "2027 纵向深化：2026 五子任务 → \"文档规范 / 系统集成 / 安全伦理\"三条线重组为三大年度任务",
            _in(0.4), _in(2.6), _in(12.5), _in(0.55), C["method"], C["edge_base"], Pt(9.5))

    # 三大任务
    tasks = [
        ("2027 任务1\n《地理空间模型上下文协议 V2.0》\n技术文档编制\n承接：交互标准化 + GAKG 协议",
         "→ 协议 V2.0 文档"),
        ("2027 任务2\n全球地理信息智能框架系统\nV2.0 研发\n承接：GAKG 注册中心+CAFE+低代码",
         "→ 框架系统 V2.0"),
        ("2027 任务3\n《地理空间人工智能应用\n伦理准则与安全规范》编制\n承接：默认安全与合规协作",
         "→ 伦理与安全规范"),
    ]
    for i, (txt, badge) in enumerate(tasks):
        x = 0.4 + i * 4.25
        add_box(slide, txt, _in(x), _in(3.4), _in(3.85), _in(1.25), C["white"], C["edge_base"], Pt(8))
        add_box(slide, badge, _in(x + 0.6), _in(4.9), _in(2.65), _in(0.45), C["arch"], C["edge_red"], Pt(9))

    # 成果出口
    add_box(slide,
            "成果出口：2 份技术文档（协议 V2.0 / 伦理与安全规范）+ 可运行框架系统 V2.0\n"
            "支撑三大应用系统升级（课题 2/3/4 V2.0），形成\"研发—应用—反馈—优化\"闭环",
            _in(0.4), _in(5.8), _in(12.5), _in(0.75), C["app"], C["edge_green"], Pt(9))


# ═══════════════════════════════════════════════════════════════
# 幻灯片 3：任务1 协议 V2.0
# ═══════════════════════════════════════════════════════════════
def slide_task_protocol(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    add_title_bar(slide, "任务1：《地理空间模型上下文协议 V2.0》技术文档编制——协议工程方法",
                  _in(0.4), _in(0.15), _in(12.5), _in(0.5), Pt(11))

    steps = [
        ("前期基础\nGeoMCP V1.0\nJSON-RPC 2.0\n双端实现", C["phase"]),
        ("需求调研\n异构系统互操作\n四层原语定义\n国际标准对标", C["core"]),
        ("协议设计\n握手/控制/数据/事件\ncontract_binding\npushdown_spec\nfederated_sources", C["core"]),
        ("SDK 实现\n跨平台通信套件\nJava/Python SDK\n契约测试对齐", C["app"]),
        ("一致性验证\nConformance Vectors\n双端 CI 独立验证\n≥13 测试用例", C["app"]),
        ("文档定稿\n协议 V2.0 技术文档\n国际标准草案\n评审发布", C["arch"]),
    ]
    for i, (txt, bg) in enumerate(steps):
        x = 0.3 + i * 2.17
        add_box(slide, txt, _in(x), _in(0.9), _in(1.87), _in(1.55), bg, C["edge_gray"], Pt(7.5))
        if i < 5:
            add_arrow(slide, _in(x + 1.87), _in(1.6), _in(0.3), _in(0.12), C["edge_gray"])

    # 关键方法
    add_box(slide,
            "关键方法：协议栈分层抽象（四层原语）· 增量演进（V1.0 兼容 V2.0）· 契约驱动开发（先向量后实现）\n"
            "原理依托：ISO/TC 211 元模型与轻量级 OGC API 融合 · JSON-RPC over WebSocket + SSE 事件推送",
            _in(0.4), _in(2.85), _in(12.5), _in(0.6), C["method"], C["edge_base"], Pt(8.5))

    # 产出
    add_box(slide,
            "2027 产出：《地理空间模型上下文协议 V2.0》技术文档 1 份（评审通过）\n"
            "配套成果：双端 SDK 一致性测试通过 · 支撑框架系统 V2.0 与课题 2/3/4 协议对接",
            _in(0.4), _in(3.85), _in(12.5), _in(0.65), C["app"], C["edge_green"], Pt(9))


# ═══════════════════════════════════════════════════════════════
# 幻灯片 4：任务2 框架系统 V2.0
# ═══════════════════════════════════════════════════════════════
def slide_task_framework(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    add_title_bar(slide, "任务2：全球地理信息智能框架系统 V2.0 研发——系统工程方法",
                  _in(0.4), _in(0.15), _in(12.5), _in(0.5), Pt(11))

    engines = [
        ("GAKG 合约注册中心 V2.0\n自动扫描流水线规模化\nPostGIS+pgvector+Neo4j\n合约满足度门控增强", C["core"]),
        ("CAFE 计算下推引擎 V2.0\nCentral+Worker 集群扩展\n多目标智能路由算法\nServerless 弹性调度", C["core"]),
        ("低代码开发环境 V2.0\nModel Builder 升级\n≥50 节点拖拽连线\n智能体挂载机制", C["app"]),
    ]
    for i, (txt, bg) in enumerate(engines):
        x = 0.4 + i * 4.25
        add_box(slide, txt, _in(x), _in(0.85), _in(3.85), _in(1.1), bg, C["edge_base"], Pt(8))

    # 集成主线
    add_box(slide, "集成主线：注册中心（资产可发现）→ 安全网关（资产可信）→ 下推引擎（资产可算）→ 低代码（资产可用）",
            _in(0.4), _in(2.25), _in(12.5), _in(0.55), C["method"], C["edge_base"], Pt(9))

    # 关键方法
    add_box(slide,
            "关键方法：组件化解耦（五大能力独立演进）· 多目标路由（数据本地性+算力余量+合规策略）\n"
            "原理依托：CAFE\"计算移动至数据端\"范式 · AI 流量预测预热（毫秒级冷启动）· 可视化状态机",
            _in(0.4), _in(3.05), _in(12.5), _in(0.55), C["phase"], C["edge_gray"], Pt(8.5))

    # 产出
    add_box(slide,
            "2027 产出：可运行框架系统 V2.0（含 Model Builder ≥50 节点流畅拖拽与智能体挂载）\n"
            "配套成果：≥5 个 GeoWorkflow 生产级运行 · 联邦数据适配器 ≥5 类 · 支撑三大应用 V2.0 集成",
            _in(0.4), _in(3.85), _in(12.5), _in(0.65), C["app"], C["edge_green"], Pt(9))


# ═══════════════════════════════════════════════════════════════
# 幻灯片 5：任务3 伦理准则与安全规范
# ═══════════════════════════════════════════════════════════════
def slide_task_safety(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    add_title_bar(slide, "任务3：《地理空间人工智能应用伦理准则与安全规范》编制——安全工程与伦理治理方法",
                  _in(0.4), _in(0.15), _in(12.5), _in(0.5), Pt(11))

    steps = [
        ("法规知识工程\n\n≥10 国测绘法规\n跨域合规本体库\nOPA 策略规则化", C["core"]),
        ("合规核验引擎\n\n协议层拦截器\n几何拓扑+敏感语义\n并发判定<50ms", C["core"]),
        ("隐私增强计算\n\n联邦学习\n安全多方计算\n差分隐私", C["app"]),
        ("伦理治理框架\n\n数据主权·算法公平\n可解释·问责机制\nGeoAI 应用边界", C["arch"]),
        ("规范编制\n\n伦理准则安全规范\n技术文档评审\n发布", C["arch"]),
    ]
    for i, (txt, bg) in enumerate(steps):
        x = 0.3 + i * 2.62
        add_box(slide, txt, _in(x), _in(0.85), _in(2.3), _in(1.6), bg, C["edge_gray"], Pt(8))
        if i < 4:
            add_arrow(slide, _in(x + 2.3), _in(1.55), _in(0.32), _in(0.12), C["edge_gray"])

    # 原理依托
    add_box(slide,
            "原理依托：零信任架构（默认拒绝、持续验证）· 隐私增强计算（数据可用不可见）· 空间数据水印（泄露追溯）\n"
            "国际对齐：UN 人工智能伦理框架 · 各国测绘法规 · 中国《新一代人工智能伦理规范》",
            _in(0.4), _in(2.85), _in(12.5), _in(0.6), C["method"], C["edge_base"], Pt(8.5))

    # 产出
    add_box(slide,
            "2027 产出：《地理空间人工智能应用伦理准则与安全规范》文档 1 份（评审通过）\n"
            "配套成果：合规引擎并发判定延迟 <50ms 达标 · 零信任安全网关全面就绪 · 支撑课题 2/3/4 合规接入",
            _in(0.4), _in(3.85), _in(12.5), _in(0.65), C["app"], C["edge_green"], Pt(9))


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)

    slide_overall(prs)
    slide_subject1(prs)
    slide_task_protocol(prs)
    slide_task_framework(prs)
    slide_task_safety(prs)

    prs.save(str(OUT))
    print(f"完成 → {OUT}")


if __name__ == "__main__":
    main()