"""生成《2027年度课题1研究方法及技术路线》docx 文档。

运行方式：
    /tmp/docxenv/bin/python gen_docx.py
输出：
    ../2027年度课题1研究方法及技术路线.docx
"""

import os, sys
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

BASE = Path(__file__).resolve().parent
IMG = BASE / "images"
OUT = BASE.parent / "2027年度课题1研究方法及技术路线.docx"

# 图片路径
IMAGES = {
    "fig1-overall.png": "图1 项目总体研究方法与技术路线图",
    "fig2-subject1.png": "图2 课题1 总体研究方法与技术路线图",
    "fig3-task-protocol.png": "图3 任务1 协议V2.0技术路线图",
    "fig4-task-framework.png": "图4 任务2 智能框架系统V2.0技术路线图",
    "fig5-task-safety.png": "图5 任务3 伦理准则与安全规范技术路线图",
}

# ── 正文内容 ──

SECTION_I_TITLE = "一、前期基础（Phase 1 成果）"

SECTION_I = (
    "截至2026年末，GeoNexus全球地理信息智能框架已完成Phase 1核心能力建设。"
    "在协议层，研发了GeoMCP V1.0地理空间交互协议（JSON-RPC 2.0），定义了geo.capabilities、"
    "geo.describe、geo.execute、geo.health四类标准方法及五类应用错误码（2000–2004），"
    "并完成了Java多语言SDK及Python geonexus-sdk的双端实现与一致性测试。"
    "在资产管理层，设计了GeoCard地理资产元模型（含spatial、temporal、bands、access、provenance等"
    "14个section的完整JSON Schema），并完成ContractValidator契约校验模块。"
    "在节点执行层，构建了GeoNode主权计算节点与CAFE Pushdown计算下推原型，初步实现了"
    "\"计算移动至数据端\"的就近执行范式。"
    "在管理面，完成了基于RuoYi/若依框架的mgbackend后台（Java控制面），"
    "与geonexus-execution-plane执行面（Python容器化部署）的\"管算协同\"双域架构，"
    "并通过GeoMCP协议建立两域之间的标准化通信管道。"
    "上述成果为2027年V2.0的深化完善奠定了工程基础与规范基础。"
)

SECTION_II_TITLE = "二、项目总体研究方法与技术路线"

SECTION_II = (
    "本项目面向构建开放、安全、普惠的全球地理信息智能框架（GeoNexus）总目标，"
    "针对异构系统互操作难、语义鸿沟、算力门槛高和安全信任缺失等核心瓶颈，"
    "总体遵循\"需求牵引—标准先行—架构设计—迭代演进—协同联动—验证闭环\"的实施路径。\n\n"
    "（1）需求牵引与问题导向。以全球南方国家落实联合国2030年可持续发展议程的紧迫需求为驱动，"
    "聚焦\"智能鸿沟\"四类核心瓶颈，将需求分解为协议标准化、合约注册、安全合规、算力调度、"
    "智能开发等关键技术问题域，指导研究方法选择与技术攻关方向。\n"
    "（2）标准先行。从通信协议栈底层定义GeoMCP交互规范与GeoCard地理资产合约数据模型，"
    "构建异构系统间的统一\"通用语言\"，为上层应用（卫星服务、SDG评估、远程学习）提供"
    "标准化的互操作协议与语义对齐底座。\n"
    "（3）\"一基座三应用\"架构设计。课题1负责全球地理信息智能框架底座的顶层设计与核心引擎研发，"
    "课题2-4分别聚焦对地观测卫星服务、可持续发展评价、远程学习三大应用方向，"
    "各课题通过GeoMCP协议、GeoCard合约注册中心、联邦适配器、可视化组件等接口紧密耦合。\n"
    "（4）迭代演进。项目采取\"原型→Alpha→Beta→业务化\"四级迭代策略。2026年为基础构建期，"
    "完成协议草案与核心引擎原型验证；2027年为深化完善期，实现协议V2.0文档化、框架系统V2.0、"
    "伦理与安全规范三大成果；2028年为集成收尾期，完成业务化上线与多国示范应用。\n"
    "（5）协同联动。课题间通过三条链路协同推进：协议链路（课题1输出GeoMCP V2与合约数据模型，"
    "课题2-4据此协议接入资产）；数据链路（课题2为课题3提供遥感指标数据基础，课题3为课题4提供案例库素材）；"
    "知识链路（课题4将各课题的知识成果整合为知识图谱，反哺课题1-3的语义检索与智能推荐）。\n"
    "（6）验证闭环。采用\"自身验证+专家评审+第三方测试+应用示范\"闭环路线。\n\n"
    "2027年度项目考核指标包括：合规引擎并发判定延迟<50ms；微调的地理空间基础模型在封闭测试集上空间推理"
    "准确率较基座模型提升≥10%；Model Builder前端支持≥50个节点的流畅拖拽连线与智能体挂载；"
    "完成\"GeoNow 2027\"大会支撑保障工作并获官方应用证明；撰写高水平学术论文3篇，申请发明专利3项，"
    "登记计算机软件著作权2项。项目总体研究方法与技术路线如图1所示。"
)

SECTION_III_TITLE = "三、课题1总体研究方法与技术路线"

SECTION_III = (
    "课题1聚焦\"全球地理信息智能框架研发与集成\"，遵循\"标准引领→安全筑基→引擎驱动→生态聚合\""
    "递进逻辑，从协议工程、安全工程、系统工程、生态工程四维方法论维度开展研究。\n\n"
    "（1）协议工程方法。以通信协议工程方法学为指导，将GeoMCP协议栈从V1.0向V2.0演进，"
    "新增contract_binding、pushdown_spec、federated_sources关键字段，"
    "从握手层、控制层、数据层、事件层四个维度定义交互原语，"
    "融合ISO/TC 211元模型标准与新一代OGC APIs，构建面向机器可读的地理资产注册底座。\n"
    "（2）安全工程方法。遵循零信任安全架构原则（\"默认拒绝、持续验证\"），"
    "将空间法律规则形式化表达为可机读的OPA策略，嵌入协议层拦截器，实现基于地理位置、"
    "数据敏感度与用户身份的动态合规核验。\n"
    "（3）系统工程方法。深度贯彻CAFE\"计算移动至数据端\"分布式计算范式，"
    "构建Central Server+N×Worker的分布式架构，通过Pushdown Protocol将序列化计算代码"
    "推送至数据所在Worker执行。采用AI流量预测预热机制实现毫秒级Serverless调度，"
    "有效消解地理计算突发流量下的冷启动延迟。\n"
    "（4）生态工程方法。以开发者社区与公共产品共享为核心，设计多维信誉评价模型"
    "（时空分辨率、位置精度、语义一致性、被引用量与实名认证状态），"
    "支持场景感知资产推荐。采用可视化状态机技术，将空间分析算法封装为标准原子能力组件，"
    "通过Model Builder实现拖拽式组装与低门槛应用集成。\n\n"
    "2026年度课题1五个子任务（数据和模型交互标准化、GAAG合约注册中心、CAFE计算下推与Worker引擎、"
    "默认安全体系与跨域合规协作、低代码智能开发环境）为上述四个方法论维度奠定了坚实的研究基础。"
    "2027年度课题1纵向深化，将五子任务按\"文档规范/系统集成/安全伦理\"三条线重组为三大年度任务。"
    "课题1总体研究方法与技术路线如图2所示。"
)

SECTION_IV_TITLE = "四、2027年度课题1子任务研究方法与技术路线"

# 4.1
TASK1_TITLE = "4.1 任务一：《地理空间模型上下文协议V2.0》技术文档编制"

TASK1_TEXT = (
    "本任务延续2026年度子任务1（数据和模型交互标准化）与子任务2（GAAG合约注册中心）"
    "的协议设计成果，采用协议工程与增量演进方法，将GeoMCP协议栈从Phase 1（V1.0）"
    "推进至Phase 2（V2.0）的正式《地理空间模型上下文协议（GeoMCP）V2.0》技术文档形态。\n\n"
    "研究方法：（1）协议栈分层抽象。以通信协议工程方法学为指导，定义握手层、控制层、"
    "数据层、事件层四层原语的精确语义、报文格式与状态转换逻辑；在V1.0基础上新增"
    "contract_binding（资产合约绑定）、pushdown_spec（计算下推规格）与federated_sources"
    "（联邦数据源声明）三个关键字段，并定义其JSON Schema约束与向后兼容策略。\n"
    "（2）契约驱动开发。先定义协议一致性测试向量（≥13个），再实现Java/Python双端SDK，"
    "通过WireMock与GeoMCPDispatcher直测两轨并行验证；参考ISO/TC 211地理信息元数据标准与"
    "OGC API系列规范，确保协议与国际标准体系的对齐能力。\n"
    "（3）增量演进机制。采用语义化版本控制（SemVer），V2.0在保持V1.0的geo.capabilities、"
    "geo.describe、geo.execute、geo.health核心方法兼容的前提下，新增支持异步Job管理、"
    "联邦查询与计算下推的扩展方法协议。双端SDK保留V1.0降级能力。\n\n"
    "本任务研究方法与技术路线如图3所示。"
)

# 4.2
TASK2_TITLE = "4.2 任务二：全球地理信息智能框架系统V2.0研发"

TASK2_TEXT = (
    "本任务延续2026年度子任务2（GAAG合约注册中心）、子任务3（CAFE计算下推与Worker引擎）"
    "与子任务5（低代码智能开发环境）的工程成果，采用系统工程与组件化方法，"
    "将三大核心引擎从原型阶段推进至生产级V2.0系统。\n\n"
    "研究方法：（1）组件化架构重构。将注册中心、CAFE引擎与低代码开发环境解耦为独立微服务，"
    "通过GeoMCP V2协议实现组件间标准化通信。注册中心升级为PostGIS+pgvector+Neo4j"
    "三重存储架构，支持语义检索与血缘图谱查询；CAFE引擎扩展为Central Server+N×Worker"
    "集群架构，引入多目标智能路由算法（以数据本地性、算力余量、合规策略为优化目标）。\n"
    "（2）Serverless弹性调度。引入基于深度学习的流量预测模型（LSTM+Attention），"
    "实现冷启动延迟毫秒级控制；采用计算预热策略与自适应并发限流，支撑海量并发地理计算任务。\n"
    "（3）Model Builder增强。将可视化状态机引擎升级为支持≥50节点流畅拖拽连线的生产级组件；"
    "开发\"智能体挂载\"机制，允许用户将预训练的遥感解译、空间查询等智能体直接嵌入工作流节点；"
    "基于预训练智能体矩阵V1.0，完成地理空间基础模型的领域微调与评估，在封闭测试集上空间推理"
    "准确率较基座模型提升≥10%，支撑智能体在专业空间分析场景中的精准挂载。\n"
    "（4）联邦数据适配器。完成不低于5类主流对地观测联邦数据源的适配器开发与标准化接入，"
    "包括Sentinel Hub（STAC API）、OpenStreetMap（Overpass API）、Landsat（USGS M2M API）等，"
    "支撑跨域数据资源的统一检索与按需调度。\n"
    "（5）集成验证。基于课题2-4的应用集成需求驱动框架系统V2.0的接口设计，通过≥5个GeoWorkflow"
    "生产级运行进行端到端验证，形成\"框架—应用\"反馈迭代循环。\n\n"
    "本任务研究方法与技术路线如图4所示。"
)

# 4.3
TASK3_TITLE = "4.3 任务三：《地理空间人工智能应用伦理准则与安全规范》编制"

TASK3_TEXT = (
    "本任务延续2026年度子任务4（默认安全体系与跨域合规协作技术）的安全架构成果，"
    "采用安全工程与伦理治理方法，将零信任合规引擎从原型推进为可发布的技术规范体系。\n\n"
    "研究方法：（1）法规知识工程。系统梳理≥10个国家与地区的测绘与地理信息法律法规，"
    "利用知识抽取技术（NER+RE+实体对齐），构建空间法律规则本体库，并通过OPA（开放策略代理）"
    "将法规约束形式化为可执行的安全策略。重点覆盖数据跨境传输、分辨率限制、敏感区域保护等核心议题。\n"
    "（2）合规引擎性能优化。在协议层嵌入拦截器，实现\"几何拓扑+敏感语义\"双重核验——"
    "几何拓扑核验基于空间索引（R-tree）快速判定操作区域与主权边界关系；敏感语义核验基于"
    "规则引擎判定数据敏感等级。优化判定路径，通过预编译策略缓存与并行判定技术，将并发判定延迟降低至<50ms。\n"
    "（3）隐私增强计算。融合联邦学习（模型梯度交换，原始数据不出域）、安全多方计算（多方联合统计，各自输入保密）"
    "与差分隐私（聚合结果添加校准噪声）三类技术，构建多层级隐私保护技术栈。"
    "研发空间数据水印算法，在栅格瓦片频域与矢量坐标微扰中嵌入可追溯用户指纹。\n"
    "（4）伦理治理框架。以\"数据主权—算法公平—可解释性—问责制\"为四根支柱，"
    "结合UN人工智能伦理框架与中国《新一代人工智能伦理规范》，定义地理空间AI应用的伦理边界、"
    "透明度要求与审计机制，最终汇编为正式技术规范文档。\n\n"
    "本任务研究方法与技术路线如图5所示。"
)


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
    return h


def add_para(doc, text, indent_first_line=True):
    p = doc.add_paragraph()
    if indent_first_line:
        p.paragraph_format.first_line_indent = Cm(0.75)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.35
    run = p.add_run(text)
    run.font.size = Pt(10.5)
    run.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
    return p


def add_image(doc, filename, caption):
    path = IMG / filename
    if not path.exists():
        print(f"  ⚠ 图片缺失: {path}", file=sys.stderr)
        return
    # 图片
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Cm(15.0))
    # 图注
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(2)
    cap.paragraph_format.space_after = Pt(10)
    cr = cap.add_run(caption)
    cr.font.size = Pt(9)
    cr.font.italic = True
    cr.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)


def main():
    print("生成 2027年度课题1研究方法及技术路线.docx ...")
    doc = Document()

    # 页面设置
    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    # 标题
    title = doc.add_heading("2027年度课题1研究方法及技术路线", level=0)
    for run in title.runs:
        run.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run("（全球地理信息智能框架研发与集成 — 补充说明）")
    sr.font.size = Pt(10)
    sr.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

    doc.add_paragraph()  # spacing

    # ── 第一节：前期基础 ──
    add_heading(doc, SECTION_I_TITLE, level=1)
    add_para(doc, SECTION_I)

    # ── 第二节：项目总体 ──
    add_heading(doc, SECTION_II_TITLE, level=1)
    add_para(doc, SECTION_II)
    add_image(doc, "fig1-overall.png", IMAGES["fig1-overall.png"])

    # ── 第三节：课题1总体 ──
    add_heading(doc, SECTION_III_TITLE, level=1)
    add_para(doc, SECTION_III)
    add_image(doc, "fig2-subject1.png", IMAGES["fig2-subject1.png"])

    # ── 第四节：三个子任务 ──
    add_heading(doc, SECTION_IV_TITLE, level=1)

    # 4.1
    add_heading(doc, TASK1_TITLE, level=2)
    add_para(doc, TASK1_TEXT)
    add_image(doc, "fig3-task-protocol.png", IMAGES["fig3-task-protocol.png"])

    # 4.2
    add_heading(doc, TASK2_TITLE, level=2)
    add_para(doc, TASK2_TEXT)
    add_image(doc, "fig4-task-framework.png", IMAGES["fig4-task-framework.png"])

    # 4.3
    add_heading(doc, TASK3_TITLE, level=2)
    add_para(doc, TASK3_TEXT)
    add_image(doc, "fig5-task-safety.png", IMAGES["fig5-task-safety.png"])

    # 保存
    doc.save(str(OUT))
    print(f"完成 → {OUT}")

if __name__ == "__main__":
    main()