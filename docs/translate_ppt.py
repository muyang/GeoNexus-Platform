#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
将 ISPRS-public goods.pptx 从英文翻译成中文，保持格式不变。
"""

from pptx import Presentation
from pptx.util import Pt
import copy

def translate_text(original):
    """翻译映射表"""
    translations = {
        # Slide 1
        "1. Global Geospatial Public Goods": "1. 全球地理空间公共产品",
        "Global geospatial public goods encompass geospatial data, knowledge, technologies and services that can be universally shared and utilized across nations and populations. As core constituents of global governance systems, these public goods serve as essential foundations for sustainable development, climate action, and humanitarian assistance—enabling evidence-based policymaking and international collaboration.": "全球地理空间公共产品包括可在各国和人群中普遍共享和使用的地理空间数据、知识、技术和服务。作为全球治理系统的核心组成部分，这些公共产品是可持续发展、气候行动和人道主义援助的重要基础——为实现循证决策和国际合作提供支撑。",
        "Foundation geospatial data": "基础地理空间数据",
        "(1) Geospatial data": "(1) 地理空间数据",
        "Spatiotemporal thematic data\nSocioeconomic and environmental data\nSpatiotemporal SDG indicators": "时空专题数据\n社会经济和环境数据\n时空可持续发展目标指标",
        "(2) Geospatial knowledge": "(2) 地理空间知识",
        "SDG spatiotemporal monitoring": "可持续发展目标时空监测",
        "UN-IGIF implementation guidance": "联合国地理信息一体化框架实施指南",
        "(3) Technologies and Services": "(3) 技术与服务",
        "Global Geospatial Intelligence Framework\nAnd Service Hubs": "全球地理空间智能框架\n及服务枢纽",
        "Crowd-powered data mining \nFusion Change Inspection\nIntelligent Crowdsourcing Update\nCollaborative Quality Inspection": "众源数据挖掘\n融合变化检测\n智能众包更新\n协同质量检验",
        "Scenario-based spatiotemporal knowledge service \nIntelligent analysis and decision-making\nKnowledge management and service": "基于场景的时空知识服务\n智能分析与决策\n知识管理与服务",

        # Slide 2
        "1.1 Global Geospatial Data": "1.1 全球地理空间数据",
        "Establish a standardized spatiotemporal fusion framework by integrating multi-modal data resources, including global/regional statistics, Earth observation big data, and large-model geospatial embeddings, to produce spatiotemporal products with high resolution and high precision.": "通过整合多模态数据资源，包括全球/区域统计数据、地球观测大数据和大模型地理空间嵌入，建立标准化的时空融合框架，以生成高分辨率和高精度的时空产品。",
        "Statistics": "统计数据",
        "Multi-modal heterogeneous geospatial data": "多模态异构地理空间数据",
        "POIs": "兴趣点",
        "Remote\nSensing": "遥感",
        "Foundation\nGeospatialData": "基础\n地理空间数据",
        "Fusion and assimilation": "融合与同化",
        "Population structure reconstruction\nSpatialization of GDP by Sector\nStatistical Downscaling": "人口结构重构\n分行业GDP空间化\n统计降尺度",
        "Methodologies": "方法论",
        "Geospatial Expression of\nPublic Services": "公共服务的\n地理空间表达",
        "Accessibility modelling\nSocial equity Evaluation": "可达性建模\n社会公平评估",
        "Consistency processing and \nmulti-scale integration": "一致性处理与\n多尺度集成",
        "Long-term temporal reconstruction\nSpatiotemporal consistency improvement": "长期时序重建\n时空一致性提升",
        "Ten spatiotemporal dataset": "十个时空数据集",
        "Social foundation dataset": "社会基础数据集",
        "Population age and sex structure\nMedical accessibility\nEducation accessibility": "人口年龄性别结构\n医疗可达性\n教育可达性",
        "Economic foundation dataset": "经济基础数据集",

        # Slide 3
        "1.2 Global Geospatial Knowledge": "1.2 全球地理空间知识",
        "Construct a global spatiotemporal knowledge graph by extracting and structuring multi-source geospatial knowledge, which enables cross-domain associations and spatiotemporal reasoning to support knowledge services.": "通过提取和结构化多源地理空间知识，构建全球时空知识图谱，实现跨领域关联和时空推理，以支撑知识服务。",
        "Global spatiotemporal knowledge graph": "全球时空知识图谱",
        "SDG indicator knowledge": "可持续发展目标指标知识",
        "Geospatial thematic knowledge": "地理空间专题知识",
        "Geospatial knowledge extraction": "地理空间知识提取",
        "SDG indicator knowledge extraction": "可持续发展目标指标知识提取",
        "Thematic knowledge extraction": "专题知识提取",
        "Knowledge graph fusion and reasoning": "知识图谱融合与推理",
        "Spatiotemporal knowledge service": "时空知识服务",
        "Indicator monitoring": "指标监测",
        "Spatiotemporal analysis": "时空分析",
        "Knowledge Q&A": "知识问答",

        # Slide 4
        "1.3 Global Geospatial Intelligence Framework and Service Hubs": "1.3 全球地理空间智能框架及服务枢纽",
        "Build a geospatial intelligence framework and service hub that integrates data, models, and computing power to provide scenario-based spatiotemporal knowledge services, supporting intelligent analysis and decision-making.": "构建一个集成数据、模型和算力的地理空间智能框架和服务枢纽，提供基于场景的时空知识服务，支撑智能分析与决策。",
        "Geospatial Intelligence Framework": "地理空间智能框架",
        "Service Hub": "服务枢纽",
        "Scenario-based knowledge service": "基于场景的知识服务",
        "Intelligent analysis and decision-making": "智能分析与决策",
        "Data resources": "数据资源",
        "Model resources": "模型资源",
        "Computing resources": "计算资源",
        "Scenario 1: Urban planning": "场景1：城市规划",
        "Scenario 2: Disaster response": "场景2：灾害应对",
        "Scenario 3: Environmental monitoring": "场景3：环境监测",
        "Spatiotemporal knowledge graph": "时空知识图谱",
        "Geospatial AI models": "地理空间人工智能模型",
        "High-performance computing": "高性能计算",

        # Slide 5
        "1.4 Technologies and Services": "1.4 技术与服务",
        "Develop key technologies including crowd-powered data mining, intelligent crowdsourcing update, and collaborative quality inspection to ensure the timeliness, accuracy, and reliability of geospatial public goods.": "开发众源数据挖掘、智能众包更新和协同质量检验等关键技术，以确保地理空间公共产品的时效性、准确性和可靠性。",
        "Crowd-powered data mining": "众源数据挖掘",
        "Intelligent Crowdsourcing Update": "智能众包更新",
        "Collaborative Quality Inspection": "协同质量检验",
        "Technology 1: Crowd-powered data mining": "技术1：众源数据挖掘",
        "Technology 2: Intelligent Crowdsourcing Update": "技术2：智能众包更新",
        "Technology 3: Collaborative Quality Inspection": "技术3：协同质量检验",
        "Multi-source data fusion": "多源数据融合",
        "Change detection": "变化检测",
        "Data quality assessment": "数据质量评估",
        "Data updating": "数据更新",
        "Quality control": "质量控制",
        "Service evaluation": "服务评估",
    }

    return translations.get(original, original)


def main():
    input_path = "/Users/mac/Repos/GeoNexus/docs/ISPRS-public goods.pptx"
    output_path = "/Users/mac/Repos/GeoNexus/docs/ISPRS-public goods_中文.pptx"

    print(f"正在读取PPT: {input_path}")
    prs = Presentation(input_path)
    print(f"PPT加载完成，共 {len(prs.slides)} 页")

    translated_count = 0
    skipped_count = 0

    for slide_idx, slide in enumerate(prs.slides):
        print(f"\n处理第 {slide_idx + 1} 页...")
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue

            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    original_text = run.text.strip()
                    if not original_text:
                        continue

                    # Try exact match
                    translated = translate_text(original_text)
                    if translated != original_text:
                        run.text = translated
                        translated_count += 1
                    else:
                        # Try multi-line match (preserve newlines)
                        lines = original_text.split('\n')
                        translated_lines = []
                        changed = False
                        for line in lines:
                            line_stripped = line.strip()
                            translated_line = translate_text(line_stripped)
                            if translated_line != line_stripped:
                                translated_lines.append(translated_line)
                                changed = True
                            else:
                                translated_lines.append(line)

                        if changed:
                            run.text = '\n'.join(translated_lines)
                            translated_count += 1
                        else:
                            skipped_count += 1

    print(f"\n翻译完成: {translated_count} 处翻译, {skipped_count} 处跳过（未匹配或无需翻译）")

    print(f"\n正在保存到: {output_path}")
    prs.save(output_path)
    print("保存完成！")

    import os
    size = os.path.getsize(output_path)
    print(f"文件大小: {size / 1024 / 1024:.2f} MB")


if __name__ == "__main__":
    main()
