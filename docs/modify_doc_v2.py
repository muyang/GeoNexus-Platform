#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
基于白皮书和原型代码，对实施方案615.docx进行第二轮修改。
修改部分使用黄色高亮显示。
"""

from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def add_highlighted_run(paragraph, text):
    """在段落末尾添加黄色高亮的文本"""
    run = paragraph.add_run(text)
    r = run._r
    rPr = r.get_or_add_rPr()
    highlight = OxmlElement('w:highlight')
    highlight.set(qn('w:val'), 'yellow')
    rPr.append(highlight)
    return run


def main():
    input_path = "/Users/mac/Repos/GeoNexus/docs/实施方案615.docx"
    output_path = "/Users/mac/Repos/GeoNexus/docs/实施方案615_v2.docx"

    print(f"正在读取文档: {input_path}")
    doc = Document(input_path)
    print(f"文档加载完成，共 {len(doc.paragraphs)} 个段落，{len(doc.tables)} 个表格")

    # ==================== 修改A: 在"研究内容"后增加总体框架设计理念与名词解释 ====================
    print("\n[修改A] 补充总体框架设计理念与名词解释...")

    # 找到"研究内容"标题后的段落（通常是段落90附近）
    for i, para in enumerate(doc.paragraphs):
        if '（四） 研究内容' in para.text and 'Heading' in str(para.style.name):
            print(f"  -> 找到'研究内容'标题在段落{i}")
            # 在下一个段落（研究内容的具体描述）后添加名词解释
            if i + 1 < len(doc.paragraphs):
                target_para = doc.paragraphs[i + 1]
                add_highlighted_run(target_para, "\n\n【新增内容——总体框架设计理念与核心名词解释】""")
                add_highlighted_run(target_para, "\n\n本项目总体架构遵循GeoNexus白皮书提出的'协议优先、联邦优先'设计理念，核心目标是构建一个开放、安全、普惠的全球地理空间智能基础设施。以下为关键术语定义：")

                add_highlighted_run(target_para, "\n\n（1）GeoNexus（全球地理空间智能枢纽）：")
                add_highlighted_run(target_para, "GeoNexus是一个联邦式GeoAI互操作框架，旨在通过开放协议连接全球地理空间数据集、AI模型、分析工作流、智能代理和计算能力。其核心特征包括：联邦优先（无中央数据仓库，节点保留主权）、AI原生（每个能力都以机器可读、AI可检索的语义描述）、协议中心化（通过开放、版本化的协议实现互操作，而非单体引擎）、能力导向（共享单元是Skill而非数据转储）、主权保护（合规约束是协议的一等公民）、开放生态（为社区贡献而设计，非单一厂商控制）、云原生和边缘就绪（从全球云租户到本地边缘节点均可运行）、人机协作（Agent提出工作流，人类审批、引导和解释）。")

                add_highlighted_run(target_para, "\n\n（2）GeoMCP（地理空间模型上下文协议）：")
                add_highlighted_run(target_para, "GeoMCP是GeoNexus的通用互操作协议。它受JSON-RPC 2.0启发，兼容RESTful，定义了所有GeoNexus交互的最小方法集。核心方法包括：discover（发现匹配自然语言或结构化查询的GeoCard）、describe（获取特定资产的完整GeoCard）、execute（在指定空间约束下调用GeoSkill）、status（监控长时间运行任务的进度）。每个请求携带空间上下文（CRS、边界框）和可选的安全/追踪凭证。协议传输无关，内部基于gRPC，外部基于RESTful HTTPS，在适当时与OGC服务兼容。GeoMCP与Anthropic的MCP协议保持兼容，同时扩展了空间上下文、地理计算语义和数据主权标识等地理空间专用原语。")

                add_highlighted_run(target_para, "\n\n（3）GeoCard（地理空间资产语义身份卡）：")
                add_highlighted_run(target_para, "GeoCard是任何GeoNexus资产的机器可读护照。它替代了临时元数据，提供标准化、AI友好的契约。每个GeoCard至少包含：身份信息（id、name、type、description）、能力描述（inputs、outputs、空间分辨率、时间覆盖）、执行信息（runtime、entrypoint、资源需求）、合规信息（licenses、数据主权规则、使用约束）、信任信息（来源、维护者信誉、验证状态）。GeoCard基于JSON Schema定义，支持YAML和JSON两种格式。")

                add_highlighted_run(target_para, "\n\n（4）GeoSkill（地理空间分析能力运行时标准）：")
                add_highlighted_run(target_para, "GeoSkill是地理空间能力的可执行体现。它将代码、环境和GeoCard打包为版本化、可测试的单元。标准结构包括：geocard.yaml（语义身份）、executor.py（核心执行逻辑）、workflow.yaml（工作流DAG定义）、environment.yaml（Conda/Docker环境）、tests/（验证套件）、README.md。GeoSkill尽可能无状态，接受结构化输入，返回标准化输出。它们可以被GeoAgent链接为有向无环图（DAG）。GeoSkill是GeoNexus能力市场的基本单元。")

                add_highlighted_run(target_para, "\n\n（5）GeoAgent（AI原生编排框架）：")
                add_highlighted_run(target_para, "GeoAgent是AI原生编排层。它接收自然语言或结构化意图，从注册表检索相关GeoCard，规划多步骤工作流，并跨GeoNode协调执行。核心职责包括：意图解析（将用户请求转换为结构化目标）、语义检索（使用嵌入和知识图查询GeoCard注册表）、工作流规划（生成优化的GeoSkill调用DAG）、执行路由（将任务分派到适当的GeoNode，尊重主权和能力约束）、结果合成（收集输出，生成地图、摘要报告和结构化数据）。GeoAgent基于LangGraph实现确定性、可审计的规划，避免无界自主循环。本项目将充分借鉴AutoGIS框架的合同驱动设计理念，将数据合同和算法合同机制融入GeoAgent编排逻辑。")

                add_highlighted_run(target_para, "\n\n（6）GeoNode（主权执行环境）：")
                add_highlighted_run(target_para, "GeoNode是GeoNexus的主权执行单元。它可以是单个服务器、Kubernetes集群或边缘设备，由数据托管方拥有和运营。GeoNode提供：本地GeoCard注册表、GeoMCP端点（用于外部发现和调用）、Skill Runtime环境（容器化或WASM沙箱）、策略引擎（OPA，用于执行本地合规规则）。核心不变式：数据保持本地化。GeoNode就地执行Skill，仅共享派生结果（如掩码、统计摘要）回请求者。本项目将OGE的时空立方体组织模型和分布式计算能力作为GeoNode的核心技术基础，利用CubeRDD和AI Cube实现GeoSkill的高效执行和分布式推理。")

                add_highlighted_run(target_para, "\n\n（7）GeoTrust（信誉与治理框架）：")
                add_highlighted_run(target_para, "在联邦生态中，信任不能假设，而必须可计算。GeoTrust基于以下因素分配动态信誉分数：来源记录（谁创建了这个GeoCard/Skill？）、社区验证和使用信号、执行审计日志、治理背书（如联合国机构验证）。GeoTrust是可选但强大的层，使请求者能够按可靠性和合规态势过滤能力。")

                print("  -> 已添加总体框架设计理念与名词解释")
                break

    # ==================== 修改B: 补充原型代码运行结果 ====================
    print("\n[修改B] 补充GeoNexus原型代码运行结果...")

    # 找到"创新性"部分，在其后添加原型设计说明
    for i, para in enumerate(doc.paragraphs):
        if '（三）项目创新性' in para.text and 'Heading' in str(para.style.name):
            if i + 1 < len(doc.paragraphs):
                target_para = doc.paragraphs[i + 1]
                add_highlighted_run(target_para, "\n\n【新增内容——GeoNexus原型系统设计与验证】")

                add_highlighted_run(target_para, "\n\n为验证GeoNexus总体架构的可行性和用户体验，本项目团队基于GeoMCP构想开发了GeoNexus原型系统（v13），并完成了功能验证。原型系统采用前后端分离架构，前端基于React + Tailwind CSS + Leaflet地图引擎，后端基于FastAPI + WebSocket实现异步通信。")

                add_highlighted_run(target_para, "\n\n（1）原型系统核心功能验证：")
                add_highlighted_run(target_para, "\n① 多语言国际化支持：原型系统支持中文、英文、法文、德文、意大利文、西班牙文、阿拉伯文共7种语言，通过数据属性驱动（data-i18n）实现界面文本的动态切换，满足不同国家和地区用户的使用需求。")
                add_highlighted_run(target_para, "\n② ChatMap交互界面：左侧为AI对话面板（Copilot），右侧为交互式地图（Map）。用户可以通过自然语言输入需求（如'分析南非洪水风险'），系统返回数据市场、工具箱、模拟器、案例库四个功能入口。")
                add_highlighted_run(target_para, "\n③ 数据市场（DataHub）：展示热门数据集（Sentinel-2 Mosaic、OpenBuildings、ERA5-Land等）、SOTA模型（Wildfire_Dynamics_v9、FloodGPT-Heavy、CropYield-LSTM等）和工具评分。")
                add_highlighted_run(target_para, "\n④ 工作流编排器（Model Builder）：支持节点拖拽式工作流构建。用户可以从左侧拖入数据源和工具节点，系统自动高亮兼容节点，一键运行生成任务DAG。")
                add_highlighted_run(target_para, "\n⑤ 主权与合规视觉反馈：当用户尝试将高敏感度数据连线给不符合主权策略的模型时，连线自动断裂并显示红色警告图标。敏感数据在地图上覆盖'GeoMCP Secure View'水印，坐标进行随机抖动处理。")
                add_highlighted_run(target_para, "\n⑥ 多语言社区论坛：模拟StackOverflow + Reddit混合体，支持技术问答和政策讨论。")

                add_highlighted_run(target_para, "\n\n（2）原型系统技术架构：")
                add_highlighted_run(target_para, "\n前端：React + Tailwind CSS + Leaflet.js（地图渲染）+ Lucide（图标库）；后端：FastAPI + WebSocket + PostgreSQL/PostGIS + Milvus（向量检索）+ MinIO（对象存储）；协议层：GeoMCP（JSON-RPC over WebSocket/SSE）。")

                add_highlighted_run(target_para, "\n\n（3）原型系统对项目实施的启示：")
                add_highlighted_run(target_para, "\n① 验证了'协议驱动、能力导向'架构的可行性，GeoMCP协议能够有效支撑异构资产的发现、调用和编排；② 多语言支持和国际化设计是服务发展中国家的必要条件；③ ChatMap交互模式有效降低了非专业用户的使用门槛；④ 主权与合规的视觉反馈机制增强了系统的安全性和可信度。这些原型验证结果为项目后续的系统研发和工程实现提供了重要参考。")

                print("  -> 已添加GeoNexus原型系统设计与验证")
                break

    # ==================== 保存文档 ====================
    print(f"\n正在保存文档到: {output_path}")
    doc.save(output_path)
    print("文档保存完成！")
    print(f"\n共完成2处新增内容，均以黄色高亮显示。")
    print(f"新增内容涵盖：总体框架设计理念与核心名词解释、GeoNexus原型系统设计与验证。")


if __name__ == "__main__":
    main()
