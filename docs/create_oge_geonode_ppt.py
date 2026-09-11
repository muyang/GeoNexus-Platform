#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
创建 OGE 作为 GeoNode 接入 GeoNexus 的 PPT
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# 创建演示文稿
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

def add_title_slide(prs, title, subtitle):
    slide_layout = prs.slide_layouts[6]  # 空白布局
    slide = prs.slides.add_slide(slide_layout)
    
    # 背景色
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(0x1A, 0x2B, 0x3C)
    
    # 标题
    title_box = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11.333), Inches(1.5))
    tf = title_box.text_frame
    tf.text = title
    p = tf.paragraphs[0]
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    p.alignment = PP_ALIGN.CENTER
    
    # 副标题
    sub_box = slide.shapes.add_textbox(Inches(1), Inches(4.2), Inches(11.333), Inches(1))
    tf = sub_box.text_frame
    tf.text = subtitle
    p = tf.paragraphs[0]
    p.font.size = Pt(22)
    p.font.color.rgb = RGBColor(0xAA, 0xCC, 0xEE)
    p.alignment = PP_ALIGN.CENTER
    
    # 底部装饰线
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.5), Inches(5.5), Inches(4.333), Pt(2))
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(0x00, 0x99, 0xCC)
    line.line.color.rgb = RGBColor(0x00, 0x99, 0xCC)
    
    return slide

def add_content_slide(prs, title, bullets):
    slide_layout = prs.slide_layouts[6]  # 空白布局
    slide = prs.slides.add_slide(slide_layout)
    
    # 背景色
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    
    # 顶部色条
    top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.15))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = RGBColor(0x00, 0x66, 0x99)
    top_bar.line.fill.background()
    
    # 标题
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(0.8))
    tf = title_box.text_frame
    tf.text = title
    p = tf.paragraphs[0]
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0x00, 0x44, 0x66)
    
    # 内容
    content_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(11.733), Inches(5.5))
    tf = content_box.text_frame
    tf.word_wrap = True
    
    for i, bullet in enumerate(bullets):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = bullet
        p.font.size = Pt(18)
        p.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        p.space_after = Pt(12)
        p.level = 0
    
    return slide

# Slide 1: 封面
add_title_slide(prs, 
    "OGE 作为 GeoNode 接入 GeoNexus\n系统解决方案", 
    "开放地球引擎与联邦化 GeoAI 互操作框架的深度融合")

# Slide 2: 总体架构设计
add_content_slide(prs, "总体架构设计：OGE-GeoNode 映射关系", [
    "• 内核-外壳双层架构：内核保留 OGE 完整能力（GeoCube + CubeRDD + AI Cube），外壳为 GeoNode 标准化封装层",
    "• 本地资源底座：OGE 硬件基础设施 + 基础数据资源（全球遥感影像、DEM、定量遥感产品、国产卫星数据）",
    "• GeoCard 注册中心：维护数据集、算子、模型的语义身份，支持向量相似度搜索和语义发现",
    "• GeoMCP 协议端点：处理联邦网络的 discover / describe / execute / status 请求",
    "• OPA 策略引擎：执行数据主权与合规策略，确保敏感数据不外流",
    "• 运行时适配器：将 GeoMCP 请求翻译为 OGEScript / OGC API，结果反向翻译为 GeoMCP 响应"
])

# Slide 3: 数据层 - GeoCube → GeoCard
add_content_slide(prs, "数据就绪层：GeoCube → GeoCard 语义映射", [
    "• 一数据集一卡、一产品一卡、一立方体一卡：GeoCube 时空数据统一赋予 GeoCard 语义身份",
    "• Identity：全球唯一标识符、名称、类型、时空覆盖范围（基于 PostGIS 索引）",
    "• Capabilities：空间分辨率、时间频率、波段清单、数据格式、时空基准（CRS / 网格 / 瓦片）",
    "• Compliance：许可类型（CC-BY / 公共领域）、数据主权级别（global / regional / national / restricted）",
    "• Trust：数据来源（卫星型号、处理算法）、社区评分、使用统计、GeoTrust 信誉评分",
    "• 向量语义索引：通过 BGE-M3 / Sentence-BERT 嵌入描述文本，支持自然语言联邦发现"
])

# Slide 4: 计算层 - CubeRDD → GeoSkill
add_content_slide(prs, "分析就绪层：CubeRDD → GeoSkill 能力封装", [
    "• 标准封装：geocard.yaml + executor.py + environment.yaml + tests/ + README.md",
    "• 算子级封装：NDVI 计算、水体提取、土地覆盖变化检测、SAR 散斑滤波等数百个算子独立封装",
    "• 工作流级封装：workflow.yaml 定义多步 GeoSkill 调用 DAG（如 Sentinel-1 洪水制图流程）",
    "• CubeRDD 弹性计算保留：自动切分时空数据、并行处理、学习型负载均衡",
    "• 数据本地化执行：原始数据始终留在 OGE-GeoNode 本地，仅返回衍生结果（掩膜、统计摘要）",
    "• 学习型时空计算：机器学习预测计算强度，自适应分配资源，异质性负载均衡提升一个量级"
])

# Slide 5: AI 层 - AI Cube → GeoSkill
add_content_slide(prs, "AI 就绪层：AI Cube → GeoSkill 智能推理集成", [
    "• AI 模型封装：基于 LuoJiaNET 的深度学习模型（洪水检测、作物分类、森林砍伐监测等）封装为 GeoSkill",
    "• AI 增强 GeoCard：模型架构、输入张量形状、输出类别、硬件需求（GPU / CUDA）、验证精度",
    "• 动态模型调度：GeoAgent 根据区域特征和 GeoTrust 评分自动选择最优模型",
    "• CPU/GPU 协同推理：大范围异质场景的高效地理 AI 推理，仅传输推理结果",
    "• 知识图谱验证：GeoTKG 验证 AI 工作流的物理合理性和时空兼容性（如光学影像不适用季风季节）",
    "• 合约驱动执行：借鉴 AutoGIS 模式，Data Agent 验证数据合约，Code Agent 生成并修复执行代码"
])

# Slide 6: 协议层 - OGC API → GeoMCP
add_content_slide(prs, "协议适配层：OGC API → GeoMCP 标准化桥接", [
    "• OGE 原生接口：OGC API-Features / Coverages / Processes + OGEScript 脚本语言",
    "• 运行时适配器三模块：Translate（请求翻译）→ Execute（引擎调用）→ Respond（结果封装）",
    "• Discover：GeoMCP 发现请求 → OGE 元数据查询 → GeoCard 列表返回",
    "• Describe：GeoMCP 描述请求 → OGE 数据集详情 → 完整 GeoCard + 输入输出模式",
    "• Execute：GeoMCP 执行请求 → OGEScript 代码生成 → Spark 作业提交 → 结果 GeoTIFF/JSON 返回",
    "• Status：任务进度监控 → 周期性状态轮询 → 进度百分比 / 中间结果 / 最终输出"
])

# Slide 7: 安全与主权
add_content_slide(prs, "安全与主权保障机制", [
    "• 零信任安全模型：默认不信任任何节点、请求或用户，每个交互必须经过认证、授权和审计",
    "• 数据主权分级：unrestricted（全球可访问）→ regional（区域授权）→ national / restricted（本地处理）",
    "• OPA 策略引擎：Rego 策略自动评估请求方身份、组织归属和 GeoTrust 评分，执行事前门禁",
    "• 容器化 + WASM 沙箱：受信任技能 Docker 容器运行，不可信技能 WASM 沙箱隔离，防止宿主机逃逸",
    "• 不可变审计日志：密码学链接的仅追加账本，记录所有请求、策略评估、执行事件和数据访问",
    "• 合规从人工检查清单到自动化门禁：《数据安全法》《网络安全法》等法规编码为机器可执行规则"
])

# Slide 8: 典型应用与实施路径
add_content_slide(prs, "典型应用场景与三阶段实施路径", [
    "• 场景一：跨国界 SDG 指标监测 — GeoCube 统一数据底座 + AI Cube 智能解译 + CubeRDD 并行计算",
    "• 场景二：灾害风险减量与应急响应 — 实时 SensorThings API 数据接入 + 边缘 GeoNode 快速部署",
    "• 场景三：气候变化与碳排放核算 — 长时序遥感 + 气候模式 + GDP 碳排放空间化模型协同分析",
    "• 第一阶段（2026）：协议适配与节点部署 — GeoCard 自动生成 + GeoMCP 端点 + OPA 策略引擎集成",
    "• 第二阶段（2027）：能力封装与联邦测试 — 数百算子 GeoSkill 化 + 亚非拉多节点联邦验证",
    "• 第三阶段（2028）：生态扩展与标准化 — 纳入全球公共产品网络 + GeoCube / GeoCard 国际标准对接"
])

# 保存
output_path = "/Users/mac/Repos/GeoNexus/docs/OGE_as_GeoNode_Solution.pptx"
prs.save(output_path)
print(f"PPT 已保存至: {output_path}")
